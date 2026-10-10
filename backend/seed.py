from __future__ import annotations

import argparse
import hashlib
import json
import re
import sqlite3
from datetime import datetime, timedelta
from pathlib import Path
from typing import Any, Iterable, Optional, Tuple
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

from .database import DEFAULT_DATABASE, connect, init_db


EASTERN = ZoneInfo("America/New_York")
DEFAULT_WORKBOOK = Path(__file__).resolve().parent / "data" / "ru_relief_demo_dataset.xlsx"


def yes(value: Any) -> int:
    return int(str(value or "").strip().lower().startswith("yes") or str(value or "").strip() == "Prototype")


def iso_datetime(value: Any) -> Optional[str]:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        result = value
    else:
        try:
            result = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
        except ValueError:
            return str(value)
    if result.tzinfo is None:
        result = result.replace(tzinfo=EASTERN)
    return result.isoformat()


def normalize_setting_key(value: str) -> str:
    if value.startswith("DEMO_NOW"):
        return "DEMO_NOW"
    return re.sub(r"[^a-zA-Z0-9_]+", "_", value).strip("_")


def setting_value(value: Any) -> tuple[str, str]:
    if isinstance(value, datetime):
        return iso_datetime(value) or "", "datetime"
    if isinstance(value, bool):
        return ("1" if value else "0"), "boolean"
    if isinstance(value, int):
        return str(value), "integer"
    if isinstance(value, float):
        return str(value), "number"
    return str(value), "text"


def rows_from_sheet(workbook, sheet_name: str, header_row: int = 1) -> list[dict[str, Any]]:
    sheet = workbook[sheet_name]
    headers = [cell.value for cell in sheet[header_row]]
    rows: list[dict[str, Any]] = []
    for values in sheet.iter_rows(min_row=header_row + 1, values_only=True):
        if all(value is None for value in values):
            continue
        record = {str(header): value for header, value in zip(headers, values) if header is not None}
        rows.append(record)
    return rows


def parse_time_minutes(value: Any) -> Optional[int]:
    if value in (None, ""):
        return None
    if hasattr(value, "hour") and hasattr(value, "minute"):
        return int(value.hour) * 60 + int(value.minute)
    text = str(value).strip()
    hour, minute = (int(piece) for piece in text.split(":")[:2])
    return hour * 60 + minute


def internal_capacity(status: str, proposed: bool) -> int:
    if not proposed:
        return 0
    if status == "Available":
        return 2
    if status == "Low stock":
        return 1
    return 0


def hash_pickup_code(code: Optional[str]) -> Tuple[Optional[str], Optional[str]]:
    if not code:
        return None, None
    return hashlib.sha256(code.encode("utf-8")).hexdigest(), code[-4:]


def clear_seeded_data(database: sqlite3.Connection) -> None:
    for table in (
        "reservations",
        "verification_reports",
        "inventory",
        "location_hours",
        "products",
        "locations",
        "health_contacts",
        "settings",
    ):
        database.execute(f"DELETE FROM {table}")


def seed_settings(database: sqlite3.Connection, workbook) -> None:
    for row in rows_from_sheet(workbook, "Config"):
        parameter = row.get("parameter")
        if not parameter or row.get("value") is None:
            continue
        value, value_type = setting_value(row["value"])
        database.execute(
            """
            INSERT INTO settings(key, value, value_type, unit, basis, status)
            VALUES (?, ?, ?, ?, ?, ?)
            """,
            (
                normalize_setting_key(str(parameter)),
                value,
                value_type,
                row.get("unit"),
                row.get("basis"),
                row.get("status"),
            ),
        )


def seed_locations(database: sqlite3.Connection, workbook) -> None:
    day_groups = {
        0: ("mon_thu_open", "mon_thu_close"),
        1: ("mon_thu_open", "mon_thu_close"),
        2: ("mon_thu_open", "mon_thu_close"),
        3: ("mon_thu_open", "mon_thu_close"),
        4: ("fri_open", "fri_close"),
        5: ("sat_open", "sat_close"),
        6: ("sun_open", "sun_close"),
    }
    for row in rows_from_sheet(workbook, "Locations"):
        location_verified = int(row["location_id"] != "L07")
        database.execute(
            """
            INSERT INTO locations(
                location_id, location_name, campus_zone, site_description,
                latitude, longitude, coordinate_note, location_type,
                meal_swipe_eligible_today, ru_express_accepted_today,
                swipe_notes, hours_basis, proposed_health_section,
                reservation_supported, staff_verification_supported,
                student_verification_supported, product_tier,
                verification_notes, location_verified
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                row["location_id"], row["location_name"], row["campus_zone"],
                row.get("site_description"), row.get("latitude"), row.get("longitude"),
                row.get("coordinate_note"), row.get("location_type"),
                yes(row.get("meal_swipe_eligible_today")), yes(row.get("ru_express_accepted_today")),
                row.get("swipe_notes"), row.get("hours_basis"), row.get("proposed_health_section"),
                yes(row.get("reservation_supported")), yes(row.get("staff_verification_supported")),
                yes(row.get("student_verification_supported")), row.get("product_tier"),
                row.get("verification_notes"), location_verified,
            ),
        )
        for day, (open_field, close_field) in day_groups.items():
            opens = parse_time_minutes(row.get(open_field))
            closes = parse_time_minutes(row.get(close_field))
            known = int(opens is not None and closes is not None)
            database.execute(
                """
                INSERT INTO location_hours(location_id, day_of_week, opens_minute, closes_minute, hours_known)
                VALUES (?, ?, ?, ?, ?)
                """,
                (row["location_id"], day, opens, closes, known),
            )


def seed_products(database: sqlite3.Connection, workbook) -> None:
    for row in rows_from_sheet(workbook, "Products"):
        database.execute(
            """
            INSERT INTO products(
                product_id, product_name, category, active_ingredient_or_type,
                simple_definition, typical_use, label_caution, product_tier,
                quantity_limit_per_transaction, requires_staff_assistance,
                retail_swipe_status, data_basis
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                row["product_id"], row["product_name"], row["category"],
                row.get("active_ingredient_or_type"), row["simple_definition"],
                row.get("typical_use"), row["label_caution"], row["product_tier"],
                row["quantity_limit_per_transaction"], yes(row.get("requires_staff_assistance")),
                row["retail_swipe_status"], row.get("data_basis"),
            ),
        )


def seed_inventory(database: sqlite3.Connection, workbook) -> None:
    for row in rows_from_sheet(workbook, "Inventory"):
        proposed = str(row.get("proposed_to_stock")) == "Yes"
        status = row["availability_status"]
        database.execute(
            """
            INSERT INTO inventory(
                inventory_id, location_id, product_id, proposed_to_stock,
                availability_status, quantity_band, reported_confidence,
                last_verified, retail_swipe_status, payment_options,
                quantity_limit, reservation_minutes, requires_staff_assistance,
                data_label, allocatable_units_internal
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                row["inventory_id"], row["location_id"], row["product_id"], int(proposed),
                status, row.get("quantity_band") or "-", row.get("reported_confidence") or "n/a",
                iso_datetime(row.get("last_verified")), row.get("retail_swipe_status"),
                row.get("payment_options"), row.get("quantity_limit"),
                int(row.get("reservation_minutes") or 75), yes(row.get("requires_staff_assistance")),
                row.get("data_label"), internal_capacity(status, proposed),
            ),
        )


def seed_reports(database: sqlite3.Connection, workbook) -> None:
    for row in rows_from_sheet(workbook, "Verification_Reports"):
        database.execute(
            """
            INSERT INTO verification_reports(
                report_id, inventory_id, reporter_role, reporter_token,
                reported_status, reported_at, geofence_distance_m,
                geofence_radius_m, geofence_result, photo_attached,
                photo_review, verification_outcome, validity_minutes, expires_at
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                row["report_id"], row["inventory_id"], row["reporter_role"],
                row["reporter_token"], row["reported_status"], iso_datetime(row["reported_at"]),
                row.get("geofence_distance_m"), row.get("geofence_radius_m"),
                row.get("geofence_result"), yes(row.get("photo_attached")),
                row.get("photo_review"), row["verification_outcome"],
                int(row["validity_minutes"]), iso_datetime(row["expires_at"]),
            ),
        )


def seed_reservations(database: sqlite3.Connection, workbook) -> None:
    for row in rows_from_sheet(workbook, "Reservations"):
        code_hash, code_last4 = hash_pickup_code(row.get("pickup_code_mock"))
        status = row["status"]
        if status == "Expired (auto)":
            status = "Active"
        database.execute(
            """
            INSERT INTO reservations(
                reservation_id, inventory_id, reserved_by_token, created_at,
                expires_at, status, pickup_method, payment_method, quantity,
                pickup_code_hash, pickup_code_last4, collected_at
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                row["reservation_id"], row["inventory_id"], row["reserved_by_token"],
                iso_datetime(row["created_at"]), iso_datetime(row["expires_at"]), status,
                row["pickup_method"], row["payment_method"], int(row["quantity"]),
                code_hash, code_last4, iso_datetime(row.get("collected_at")),
            ),
        )


def seed_health_contacts(database: sqlite3.Connection, workbook) -> None:
    for row in rows_from_sheet(workbook, "Health_Escalation"):
        database.execute(
            """
            INSERT INTO health_contacts(
                contact_id, name, type, phone_or_contact, when_to_use, source, verified
            ) VALUES (?, ?, ?, ?, ?, ?, ?)
            """,
            (
                row["contact_id"], row["name"], row["type"], row["phone_or_contact"],
                row["when_to_use"], row.get("source"), yes(row.get("verified")),
            ),
        )


def seed_database(workbook_path: str | Path, database_path: str | Path = DEFAULT_DATABASE) -> dict[str, int]:
    workbook = load_workbook(workbook_path, data_only=True, read_only=True)
    database = connect(database_path)
    try:
        init_db(database)
        clear_seeded_data(database)
        seed_settings(database, workbook)
        seed_locations(database, workbook)
        seed_products(database, workbook)
        seed_inventory(database, workbook)
        seed_reports(database, workbook)
        seed_reservations(database, workbook)
        seed_health_contacts(database, workbook)
        database.commit()
        return {
            "locations": database.execute("SELECT COUNT(*) FROM locations").fetchone()[0],
            "products": database.execute("SELECT COUNT(*) FROM products").fetchone()[0],
            "inventory": database.execute("SELECT COUNT(*) FROM inventory").fetchone()[0],
            "proposedInventory": database.execute(
                "SELECT COUNT(*) FROM inventory WHERE proposed_to_stock = 1"
            ).fetchone()[0],
            "reports": database.execute("SELECT COUNT(*) FROM verification_reports").fetchone()[0],
            "reservations": database.execute("SELECT COUNT(*) FROM reservations").fetchone()[0],
            "healthContacts": database.execute("SELECT COUNT(*) FROM health_contacts").fetchone()[0],
        }
    finally:
        workbook.close()
        database.close()


def main() -> None:
    parser = argparse.ArgumentParser(description="Seed the RU Relief SQLite database.")
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--database", type=Path, default=DEFAULT_DATABASE)
    args = parser.parse_args()
    counts = seed_database(args.workbook, args.database)
    print(json.dumps(counts, indent=2))


if __name__ == "__main__":
    main()
